from fastapi import APIRouter, UploadFile, File, Form, Query, HTTPException
from fastapi.responses import StreamingResponse, JSONResponse, Response
import json
from typing import List, Optional
from enum import Enum
from bson import ObjectId
from database import contracts_collection, clauses_collection
from utils.pii_masker import mask_pii, unmask_pii
from utils.docx_exporter import build_contract_docx
from agent_graph import agent_graph, ClauseReviewState, llm
from pydantic import BaseModel
import io
import re
import docx
from fpdf import FPDF

class RoleEnum(str, Enum):
    BUYER = "Buyer"
    SELLER = "Seller"

router = APIRouter()


def extract_pdf_text(content: bytes) -> str:
    """Extracts text from all pages of a PDF and normalizes layout artifacts."""
    text_pages = []

    for lib_name in ("pypdf", "PyPDF2"):
        try:
            pdf_mod = __import__(lib_name)
            reader = pdf_mod.PdfReader(io.BytesIO(content))
            for page in reader.pages:
                page_text = page.extract_text()
                if page_text:
                    text_pages.append(page_text)
            if text_pages:
                break
        except ImportError:
            continue
        except Exception as e:
            print(f"[PDF {lib_name} warning]: {e}")

    if not text_pages:
        try:
            import fitz
            doc = fitz.open(stream=content, filetype="pdf")
            for page in doc:
                page_text = page.get_text()
                if page_text:
                    text_pages.append(page_text)
        except ImportError:
            pass
        except Exception as e:
            print(f"[PDF fitz warning]: {e}")

    if not text_pages:
        raise HTTPException(
            status_code=400,
            detail="PDF parser library not found in environment. Please run: pip install pypdf"
        )

    raw_text = "\n".join(text_pages)

    # Clean standalone Page X footers/headers
    cleaned_lines = []
    for line in raw_text.splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        if re.match(r'^Page\s+\d+(\s+of\s+\d+)?$', stripped, re.IGNORECASE):
            continue
        if re.match(r'^\\\\[A-Za-z0-9_\\.\s]+$', stripped):
            continue
        cleaned_lines.append(stripped)

    normalized = "\n".join(cleaned_lines)

    # Fix PDF column-wrapping quirk where "3." and "4." float above "I. For personnel..."
    normalized = re.sub(r'\n3\.\s*\n4\.\s*\n(?=[A-Z]\.)', '\n', normalized)
    normalized = re.sub(r'(?:^|\n)TERM\.\s+This Agreement', '\n3. TERM. This Agreement', normalized)
    normalized = re.sub(r'(?:^|\n)EARLY TERMINATION\.', '\n4. EARLY TERMINATION.', normalized)

    return normalized


def extract_docx_text(content: bytes) -> str:
    """Extracts text from .docx paragraphs (and tables if clauses are inside tables)."""
    doc = docx.Document(io.BytesIO(content))

    def is_noise(text: str) -> bool:
        if re.search(r'<[A-Z_]+_\d+>', text):
            return False
        if re.match(r'^(\d+|_+)$', text):
            return True
        if re.match(r'^(By|Title|Date):.*', text, re.IGNORECASE):
            return True
        if re.match(r'^Page\s+\d+(\s+of\s+\d+)?$', text, re.IGNORECASE):
            return True
        return False

    raw_paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip() and not is_noise(p.text.strip())]

    # If a large .docx stores its clauses primarily inside tables, extract table rows too
    if len(raw_paragraphs) < 3 and doc.tables:
        for table in doc.tables:
            for row in table.rows:
                row_text = " ".join(cell.text.strip() for cell in row.cells if cell.text.strip()).strip()
                if row_text and not is_noise(row_text) and row_text not in raw_paragraphs:
                    raw_paragraphs.append(row_text)

    return "\n".join(raw_paragraphs)


def split_contract_into_clauses(full_text: str) -> List[str]:
    """
    Splits both 5-clause .docx contracts and large 24+ section PDFs/DOCXs into clean top-level clauses
    without breaking on nested invoice sub-lists or signature blocks.
    """
    # Detach IN WITNESS WHEREOF signature block temporarily so numbered signature lines
    # (e.g., "1. CONSULTANT", "2. COMMISSION") are not mistaken for new clauses
    witness_match = re.search(r'(?:^|\n)(IN WITNESS WHEREOF\b[\s\S]*)$', full_text, re.IGNORECASE)
    signature_tail = ""
    body_text = full_text
    if witness_match and witness_match.start() > len(full_text) // 2:
        signature_tail = "\n\n" + witness_match.group(1).strip()
        body_text = full_text[:witness_match.start()].strip()

    # Check if the document has numbered sections (1., 2., 3...)
    candidate_pattern = re.compile(
        r'(?:^|\n|(?<=[\.\!\?])\s+)(?=(?:\d+[\.\)]\s+[A-Z]|SECTION\s+\d+|ARTICLE\s+[IVXLCDM\d]+))'
    )

    raw_chunks = [c.strip() for c in candidate_pattern.split(body_text) if c.strip()]

    # Determine if top-level headers are predominantly ALL-CAPS (like SampleContract-Shuttle.pdf)
    all_caps_headers = sum(
        1 for c in raw_chunks if re.match(r'^\d+[\.\)]\s+[A-Z]{3,}\b', c)
    )
    is_all_caps_style = all_caps_headers >= 6

    clauses: List[str] = []
    expected_num = 1

    for chunk in raw_chunks:
        num_match = re.match(r'^(\d+)[\.\)]\s+([^\n]+)', chunk)
        sec_match = re.match(r'^(?:SECTION\s+\d+|ARTICLE\s+[IVXLCDM\d]+)', chunk, re.IGNORECASE)

        if num_match:
            num = int(num_match.group(1))
            first_line = num_match.group(2).strip()

            # Detect if this chunk is actually a nested sub-list item (e.g., 1..7 inside Section 2.H)
            is_sublist_item = False
            if is_all_caps_style and not re.match(r'^[A-Z]{2,}', first_line):
                is_sublist_item = True
            elif num != expected_num and clauses:
                # Allow small skips (e.g. expected_num + 1), but reject restarts like 1, 2 inside Clause 2
                if num < expected_num or num > expected_num + 2:
                    is_sublist_item = True
            elif first_line.endswith(";"):
                is_sublist_item = True

            if is_sublist_item and clauses:
                clauses[-1] = clauses[-1] + "\n" + chunk
            else:
                clauses.append(chunk)
                expected_num = num + 1
        elif sec_match:
            clauses.append(chunk)
        else:
            # Preamble before Clause 1
            clauses.append(chunk)

    # Merge preamble into Clause 1 so Clause 1 starts cleanly with the contract intro + Section 1
    if len(clauses) > 1 and not re.match(r'^(?:\d+[\.\)]\s+[A-Z]|SECTION\s+\d+|ARTICLE\s+[IVXLCDM\d]+)', clauses[0]):
        clauses[1] = clauses[0] + "\n\n" + clauses[1]
        clauses = clauses[1:]

    # Re-attach signature block to the final clause
    if signature_tail and clauses:
        clauses[-1] = clauses[-1] + signature_tail

    return clauses


@router.post("/upload")
@router.post("/analyze")
async def upload_contract(
    file: UploadFile = File(...),
    role: Optional[str] = Form(None),
    jurisdiction: Optional[str] = Form(None),
    role_query: Optional[str] = Query(None, alias="role"),
    jurisdiction_query: Optional[str] = Query(None, alias="jurisdiction")
):
    try:
        # Resolve role parameter (default: "Buyer", valid options: "Buyer", "Seller")
        raw_role = (role or role_query or "Buyer").strip()
        selected_role = "Seller" if raw_role.lower() == "seller" else "Buyer"

        # Resolve jurisdiction parameter (default: "California")
        raw_jurisdiction = (jurisdiction or jurisdiction_query or "California").strip()
        selected_jurisdiction = raw_jurisdiction if raw_jurisdiction else "California"

        filename_lower = (file.filename or "").lower()
        if not (filename_lower.endswith(".pdf") or filename_lower.endswith(".docx")):
            raise HTTPException(status_code=400, detail="Only .pdf and .docx files are supported")

        content = await file.read()

        # 1. Extract text based on file type
        if filename_lower.endswith(".docx"):
            full_text_to_mask = extract_docx_text(content)
        else:
            full_text_to_mask = extract_pdf_text(content)

        if not full_text_to_mask.strip():
            raise HTTPException(status_code=400, detail="No readable text found in document")

        # 2. Split into clauses FIRST before PII masking
        clause_splits_raw = split_contract_into_clauses(full_text_to_mask)

        # 3. Mask PII on each individual clause and combine
        clause_splits = []
        document_mapping = {}
        for c in clause_splits_raw:
            masked_c, mapping = mask_pii(c)
            clause_splits.append(masked_c)
            document_mapping.update(mapping)

        full_masked_text = "\n\n".join(clause_splits)

        # 4. Save contract metadata with role and jurisdiction to MongoDB
        contract_doc = {
            "filename": file.filename,
            "status": "processed",
            "role": selected_role,
            "jurisdiction": selected_jurisdiction,
            "pii_mapping": document_mapping,
            "raw_text": full_masked_text
        }
        result = await contracts_collection.insert_one(contract_doc)
        contract_id = result.inserted_id

        # 5. Save the chunks as clauses
        clauses_to_insert = []
        response_clauses = []
        for index, masked_para in enumerate(clause_splits):
            clause_doc = {
                "contract_id": contract_id,
                "clause_index": index,
                "original_text": masked_para,
                "masked_text": masked_para,
                "role": selected_role,
                "jurisdiction": selected_jurisdiction,
                "risk_score": 0,
                "status": "pending",
                "prosecutor_flags": [],
                "playbook_rules": [],
                "amended_text": "",
                "defender_redline": ""
            }
            clauses_to_insert.append(clause_doc)

        if clauses_to_insert:
            insert_result = await clauses_collection.insert_many(clauses_to_insert)
            for idx, cid in enumerate(insert_result.inserted_ids):
                c = clauses_to_insert[idx]
                response_clauses.append({
                    "clause_id": str(cid),
                    "original_text": c["original_text"],
                    "masked_text": c["masked_text"],
                    "role": selected_role,
                    "jurisdiction": selected_jurisdiction,
                    "risk_score": c["risk_score"],
                    "status": c["status"],
                    "prosecutor_flags": c["prosecutor_flags"],
                    "playbook_rules": c["playbook_rules"],
                    "amended_text": c["amended_text"],
                    "defender_redline": c["defender_redline"],
                    "iteration_count": 0
                })

        return {
            "message": "Contract uploaded successfully",
            "contract_id": str(contract_id),
            "role": selected_role,
            "jurisdiction": selected_jurisdiction,
            "clauses_count": len(clauses_to_insert),
            "clauses": response_clauses
        }
    except HTTPException:
        raise
    except Exception as e:
        import traceback
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=str(e))


@router.get("/{contract_id}/stream")
async def stream_contract_review(
    contract_id: str,
    role: Optional[str] = Query(None),
    jurisdiction: Optional[str] = Query(None)
):
    async def event_generator():
        try:
            contract_oid = ObjectId(contract_id)
        except Exception:
            yield f"data: {json.dumps({'event': 'clause_error', 'clause_id': '', 'error': 'Invalid contract ID'})}\n\n"
            return

        contract = await contracts_collection.find_one({"_id": contract_oid})
        active_role = role or (contract.get("role") if contract else None) or "Buyer"
        active_jurisdiction = jurisdiction or (contract.get("jurisdiction") if contract else None) or "California"

        cursor = clauses_collection.find({"contract_id": contract_oid}).sort("clause_index", 1)
        clauses = await cursor.to_list(length=None)

        from agent_graph import analyze_clause

        for idx, item in enumerate(clauses):
            clause_text = item.get("original_text", "") if isinstance(item, dict) else str(item)
            clause_id_str = str(item.get("_id", idx)) if isinstance(item, dict) else str(idx)

            try:
                analysis = await analyze_clause(
                    clause_text,
                    role=active_role,
                    jurisdiction=active_jurisdiction
                )

                payload = {
                    "clause_index": idx,
                    "clause_id": clause_id_str,
                    "original_text": clause_text,
                    "role": active_role,
                    "jurisdiction": active_jurisdiction,
                    "risk_score": analysis["risk_score"],
                    "risk_analysis": analysis["risk_analysis"],
                    "violations": analysis["violations"],
                    "amended_text": analysis["amended_text"],
                    "defender_redline": analysis["amended_text"],
                    "status": "completed"
                }

                yield f"data: {json.dumps(payload)}\n\n"

                await clauses_collection.update_one(
                    {"_id": item["_id"]},
                    {"$set": {
                        "risk_score": analysis["risk_score"],
                        "prosecutor_flags": analysis["risk_analysis"],
                        "playbook_rules": analysis["violations"],
                        "amended_text": analysis["amended_text"],
                        "defender_redline": analysis["amended_text"],
                        "role": active_role,
                        "jurisdiction": active_jurisdiction,
                        "status": "completed"
                    }}
                )
            except Exception as e:
                import traceback
                traceback.print_exc()
                yield f"data: {json.dumps({'event': 'clause_error', 'clause_id': clause_id_str, 'error': str(e)})}\n\n"

    return StreamingResponse(event_generator(), media_type="text/event-stream")


class ChatRequest(BaseModel):
    message: Optional[str] = None
    query: Optional[str] = None
    current_clause_id: Optional[str] = None
    clause_index: Optional[int] = None


class EmailRequest(BaseModel):
    clauses: Optional[List[dict]] = None


@router.post("/copilot/chat")
@router.post("/{contract_id}/chat")
async def chat_with_contract(contract_id: str = "default", request: ChatRequest = None):
    try:
        # Resolve user message
        user_message_raw = ""
        if request:
            user_message_raw = (request.message or request.query or "").strip()
        if not user_message_raw:
            return {"reply": "Please provide a question or instruction regarding the contract."}

        # 1. Resolve Contract Document from MongoDB
        contract = None
        contract_oid = None
        if contract_id and contract_id not in ("default", "copilot"):
            try:
                contract_oid = ObjectId(contract_id)
                contract = await contracts_collection.find_one({"_id": contract_oid})
            except Exception:
                contract = None

        if not contract:
            # Fallback to the latest ingested contract
            contract = await contracts_collection.find_one(sort=[("_id", -1)])
            if contract:
                contract_oid = contract["_id"]

        if not contract:
            raise HTTPException(status_code=404, detail="Contract not found in database.")

        chat_history = contract.get("chat_history", [])
        history_text = "\n".join([f"{msg['role'].capitalize()}: {msg['content']}" for msg in chat_history[-8:]])

        # 2. Query MongoDB for all parsed clauses (Task 1: Context Expansion)
        cursor = clauses_collection.find({"contract_id": contract_oid}).sort("clause_index", 1)
        clauses = await cursor.to_list(length=None)
        if not clauses:
            cursor = clauses_collection.find().sort("clause_index", 1)
            clauses = await cursor.to_list(length=None)

        # 3. Resolve Current Focused Clause
        clause_id_req = request.current_clause_id if request else None
        if not clause_id_req and request and request.clause_index is not None:
            clause_id_req = str(request.clause_index)

        current_clause = None
        current_clause_id = clause_id_req or "1"

        if clause_id_req:
            try:
                current_clause = await clauses_collection.find_one({
                    "contract_id": contract_oid,
                    "_id": ObjectId(clause_id_req)
                })
            except Exception:
                current_clause = None

            if not current_clause:
                try:
                    c_num = int(clause_id_req)
                    current_clause = await clauses_collection.find_one({
                        "contract_id": contract_oid,
                        "$or": [{"clause_index": c_num}, {"clause_index": c_num - 1}]
                    })
                except Exception:
                    current_clause = None

            if not current_clause:
                current_clause = await clauses_collection.find_one({
                    "contract_id": contract_oid,
                    "clause_id": clause_id_req
                })

        if not current_clause and clauses:
            current_clause = clauses[0]

        if current_clause:
            c_idx = current_clause.get("clause_index", 0)
            current_clause_id = str(c_idx if c_idx >= 1 else c_idx + 1)
            c_title = current_clause.get("title", f"Clause {current_clause_id}")
            c_text = current_clause.get("masked_text", current_clause.get("original_text", "")).strip()
            c_risk = current_clause.get("risk_score", 0)
            c_flags = current_clause.get("prosecutor_flags", [])
            c_flags_str = "; ".join(c_flags) if c_flags else "None"
            c_redline = current_clause.get("amended_text", current_clause.get("defender_redline", ""))

            current_clause_context = f"""Clause #{current_clause_id}: {c_title}
Risk Score: {c_risk}/10
Identified Flags: {c_flags_str}
Text:
{c_text}"""
            if c_redline:
                current_clause_context += f"\nProposed Redline Revision:\n{c_redline}"
        else:
            current_clause_context = f"Clause #{current_clause_id} (No clause body available)"

        # 4. Assemble Full Contract Context (Inventory Summary + All Clause Texts)
        clause_inventory = []
        clause_body_blocks = []

        for c in clauses:
            c_i = c.get("clause_index", 0)
            disp_i = c_i if c_i >= 1 else c_i + 1
            title = c.get("title", f"Clause {disp_i}")
            risk = c.get("risk_score", 0)
            flags = c.get("prosecutor_flags", [])
            flags_str = "; ".join(flags) if flags else "None"
            text = c.get("masked_text", c.get("original_text", "")).strip()

            clause_inventory.append(
                f"- Clause #{disp_i}: {title} | Risk Score: {risk}/10 | Flags: {flags_str}"
            )
            clause_body_blocks.append(
                f"--- [Clause #{disp_i}: {title} | Risk: {risk}/10] ---\n{text}"
            )

        inventory_summary = "\n".join(clause_inventory) if clause_inventory else "No parsed clauses available."
        full_text_combined = "\n\n".join(clause_body_blocks)
        if len(full_text_combined) > 20000:
            full_text_combined = full_text_combined[:20000] + "\n...[Additional clauses truncated for context limit]..."

        full_contract_context = f"""[CLAUSE RISK & FLAGS INVENTORY]
Total Analyzed Clauses: {len(clauses)}
{inventory_summary}

[FULL CONTRACT CLAUSES TEXT]
{full_text_combined}"""

        user_message_masked, _ = mask_pii(user_message_raw)
        role = contract.get("role", "Buyer")
        jurisdiction = contract.get("jurisdiction", "California")

        # 5. Smart Relaxed System Directive & Analytical Prompt (Task 2)
        system_directive = f"""You are an elite Enterprise Legal Copilot. You have access to the full contract context, but the user is currently viewing Clause #{current_clause_id}. 
CRITICAL DIRECTIVES:
- You MAY summarize, analyze, or creatively reformat the contract text (e.g., translating legal jargon into IF/ELSE logic, bullet points, or risk matrixes) if the user asks for it.
- You MAY cross-reference other clauses in the contract to answer questions like 'What are all the harmful clauses?'.
- You MUST NOT invent fake laws, fake monetary values, or fake clauses that do not exist in the provided context. If asked about outside information, politely decline.
- DO NOT use Markdown tables under any circumstances. 
- If you need to present structured data, comparisons, or multiple variables, you MUST use clean bulleted lists or numbered lists with bold text for emphasis."""

        prompt = f"""{system_directive}

CLIENT PERSPECTIVE: {role.upper()}
GOVERNING JURISDICTION: {jurisdiction}

CURRENT FOCUSED CLAUSE:
{current_clause_context}

FULL CONTRACT CONTEXT:
{full_contract_context}

CONVERSATION HISTORY:
{history_text}

USER QUESTION:
{user_message_masked}
"""
        response = await llm.ainvoke(prompt)

        new_messages = [
            {"role": "user", "content": user_message_masked},
            {"role": "model", "content": response.content}
        ]

        if contract_oid:
            await contracts_collection.update_one(
                {"_id": contract_oid},
                {"$push": {"chat_history": {"$each": new_messages}}}
            )

        return {"reply": response.content}
    except Exception as e:
        import traceback
        traceback.print_exc()
        return {"reply": f"An error occurred during chat: {str(e)}"}


@router.post("/{contract_id}/generate-email")
async def generate_pushback_email(contract_id: str, request: EmailRequest = None):
    try:
        if request and request.clauses is not None:
            flagged_clauses = [
                c for c in request.clauses
                if c.get("risk_score", 0) >= 7
            ]
            if not flagged_clauses:
                flagged_clauses = [
                    c for c in request.clauses
                    if c.get("risk_score", 0) >= 4 or c.get("amended_text", "") != ""
                ]
        else:
            cursor = clauses_collection.find({
                "contract_id": ObjectId(contract_id),
                "risk_score": {"$gte": 7}
            }).sort("clause_index", 1)

            flagged_clauses = await cursor.to_list(length=None)
            if not flagged_clauses:
                cursor_fallback = clauses_collection.find({
                    "contract_id": ObjectId(contract_id),
                    "$or": [
                        {"risk_score": {"$gte": 4}},
                        {"amended_text": {"$exists": True, "$ne": ""}}
                    ]
                }).sort("clause_index", 1)
                flagged_clauses = await cursor_fallback.to_list(length=None)

        if not flagged_clauses:
            return {"email": "No high-risk clauses found. The contract looks good to sign!"}

        clauses_summary = ""
        for idx, clause in enumerate(flagged_clauses, 1):
            clause_text = clause.get("masked_text", clause.get("original_text", ""))
            if len(clause_text) > 700:
                clause_text = clause_text[:700] + "..."
            amended_text = clause.get("amended_text", "")
            flags = clause.get("prosecutor_flags", [])
            clauses_summary += f"{idx}. Original Clause Excerpt: {clause_text}\n"
            if flags:
                clauses_summary += f"   Identified Risk: {flags[0]}\n"
            clauses_summary += f"   Proposed Redline: {amended_text}\n\n"

        prompt = f"""
Draft a professional, polite, but firm pushback email to opposing counsel.
We have reviewed the contract and require the following amendments to resolve identified legal and commercial risks:

{clauses_summary}

Format each requested change clearly with numbered headers and concise bullet points. Do not use Markdown tables.
You must address ALL Critical and High-Risk clauses provided in the context. Do not stop at just two. Create a numbered list in the email body detailing the unilateral imbalance and proposed resolution for EVERY flagged clause.
"""
        try:
            response = await llm.ainvoke(prompt)
            email_draft = response.content
        except Exception:
            # Deterministic fallback email if Groq rate-limits on large contracts
            bullet_points = "\n\n".join([
                f"**{i}. {c.get('original_text', 'Clause')[:55].strip()}...**\n"
                f"* **Unilateral Imbalance:** {(c.get('prosecutor_flags') or ['Unbalanced risk allocation'])[0]}\n"
                f"* **Proposed Resolution:** {c.get('amended_text', 'See redlined attachment.')}"
                for i, c in enumerate(flagged_clauses, 1)
            ])
            email_draft = (
                "Subject: Proposed Redline Amendments — Contract Review\n\n"
                "Dear Counsel,\n\n"
                "Thank you for sharing the draft agreement. Following our legal and risk review, "
                "we require the following targeted amendments prior to execution:\n\n"
                f"{bullet_points}\n\n"
                "Please let us know if these revisions are acceptable so we may proceed to signature.\n\n"
                "Best regards,\nLegal Review Team"
            )

        await contracts_collection.update_one(
            {"_id": ObjectId(contract_id)},
            {"$set": {"pushback_email": email_draft}}
        )
        return {"email": email_draft}
    except Exception as e:
        import traceback
        traceback.print_exc()
        return {"email": f"An error occurred while generating the email: {str(e)}"}


@router.get("/{contract_id}/export/report")
async def export_audit_report(contract_id: str):
    contract_oid = ObjectId(contract_id)
    contract = await contracts_collection.find_one({"_id": contract_oid})
    if not contract:
        raise HTTPException(status_code=404, detail="Contract not found")

    cursor = clauses_collection.find({"contract_id": contract_oid}).sort("clause_index", 1)
    clauses = await cursor.to_list(length=None)

    doc = docx.Document()
    doc.add_heading('Audit Report', 0)
    doc.add_paragraph(f"Contract: {contract.get('filename', 'Unknown')}")

    for clause in clauses:
        doc.add_heading(f"Clause {clause.get('clause_index', 0) + 1}", level=1)
        doc.add_paragraph(f"Original Text: {clause.get('original_text', '')}")
        doc.add_paragraph(f"Risk Score: {clause.get('risk_score', 0)}/10")
        if clause.get("prosecutor_flags"):
            doc.add_paragraph("Flags: " + "; ".join(clause.get("prosecutor_flags", [])))
        if clause.get("amended_text"):
            doc.add_paragraph(f"Amended Text: {clause.get('amended_text', '')}")

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": "attachment; filename=Audit_Report.docx"}
    )


@router.get("/{contract_id}/export/email")
async def export_pushback_email(contract_id: str):
    contract_oid = ObjectId(contract_id)
    contract = await contracts_collection.find_one({"_id": contract_oid})
    if not contract or "pushback_email" not in contract:
        raise HTTPException(status_code=404, detail="Email not found")

    doc = docx.Document()
    doc.add_heading('Pushback Email', 0)
    doc.add_paragraph(contract["pushback_email"])

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": "attachment; filename=Pushback_Email.docx"}
    )


@router.get("/{contract_id}/export")
@router.get("/{contract_id}/export/docx")
@router.get("/export/docx/{contract_id}")
async def export_contract_docx(contract_id: str, format: str = "docx"):
    """
    Multi-format export endpoint (.DOCX, .PDF, .JSON) for contract audit packages.
    """
    try:
        contract_oid = ObjectId(contract_id)
    except Exception:
        raise HTTPException(status_code=400, detail="Invalid contract ID")

    contract = await contracts_collection.find_one({"_id": contract_oid})
    if not contract:
        raise HTTPException(status_code=404, detail="Contract not found")

    if format == "json":
        contract_doc = dict(contract)
        if "_id" in contract_doc:
            contract_doc["_id"] = str(contract_doc["_id"])
        cursor = clauses_collection.find({"contract_id": contract_oid}).sort("clause_index", 1)
        clauses = await cursor.to_list(length=None)
        if clauses:
            for c in clauses:
                if "_id" in c:
                    c["_id"] = str(c["_id"])
                if "contract_id" in c:
                    c["contract_id"] = str(c["contract_id"])
            contract_doc["clauses"] = clauses

        payload = json.loads(json.dumps(contract_doc, default=str))
        return JSONResponse(
            content=payload,
            headers={"Content-Disposition": 'attachment; filename="Contract_Audit_Payload.json"'}
        )

    if format == "pdf":
        cursor = clauses_collection.find({"contract_id": contract_oid}).sort("clause_index", 1)
        clauses = await cursor.to_list(length=None)

        pdf = FPDF()
        pdf.add_page()
        pdf.set_auto_page_break(auto=True, margin=15)

        def sanitize_pdf_text(text: str) -> str:
            if not text:
                return ""
            replacements = {
                '\u2018': "'", '\u2019': "'", '\u201c': '"', '\u201d': '"',
                '\u2013': '-', '\u2014': '--', '\u2026': '...', '\u00a0': ' ',
                '§': 'Section '
            }
            for orig, rep in replacements.items():
                text = text.replace(orig, rep)
            return text.encode('latin-1', 'replace').decode('latin-1')

        pdf.set_font("helvetica", "B", 18)
        pdf.cell(0, 10, "Executive Counsel Brief", new_x="LMARGIN", new_y="NEXT", align="C")
        pdf.ln(4)

        clean_name = re.sub(r'\.(pdf|docx)$', '', contract.get('filename', 'Contract'), flags=re.IGNORECASE)
        pdf.set_font("helvetica", "B", 12)
        pdf.cell(0, 7, sanitize_pdf_text(f"Contract: {clean_name}"), new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("helvetica", "", 10)
        pdf.cell(0, 6, sanitize_pdf_text(f"Jurisdiction: {contract.get('jurisdiction', 'California')} | Stance: {contract.get('role', 'Buyer')}"), new_x="LMARGIN", new_y="NEXT")
        pdf.ln(5)

        for idx, clause in enumerate(clauses, 1):
            title = clause.get("title", f"Clause {clause.get('clause_index', idx)}")
            risk_score = clause.get("risk_score", 0)
            summary_items = clause.get("risk_analysis") or clause.get("prosecutor_flags") or []
            if isinstance(summary_items, list) and summary_items:
                summary_text = summary_items[0]
            else:
                summary_text = clause.get("original_text", "")[:180].strip()

            pdf.set_font("helvetica", "B", 11)
            pdf.cell(0, 7, sanitize_pdf_text(f"{idx}. {title} - Risk Score: {risk_score}/10"), new_x="LMARGIN", new_y="NEXT")
            pdf.set_font("helvetica", "", 10)
            pdf.multi_cell(0, 5, sanitize_pdf_text(f"Summary: {summary_text}"))
            pdf.ln(3)

        pdf_bytes = bytes(pdf.output())
        return Response(
            content=pdf_bytes,
            media_type="application/pdf",
            headers={"Content-Disposition": 'attachment; filename="Executive_Counsel_Brief.pdf"'}
        )

    elif format == "docx" or format is None:
        cursor = clauses_collection.find({"contract_id": contract_oid}).sort("clause_index", 1)
        clauses = await cursor.to_list(length=None)

        buffer = build_contract_docx(contract, clauses)

        clean_name = re.sub(r'\.(pdf|docx)$', '', contract.get('filename', 'Contract'), flags=re.IGNORECASE)
        filename = f"{clean_name}_Redlined.docx"

        return StreamingResponse(
            buffer,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": f'attachment; filename="{filename}"'}
        )
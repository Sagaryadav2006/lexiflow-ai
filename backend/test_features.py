import asyncio
import io
import docx
from bson import ObjectId
from agent_graph import check_deterministic_guardrails, ClauseReviewState
from utils.docx_exporter import build_contract_docx


def test_california_non_compete_guardrail():
    print("--- Testing California Non-Compete Guardrail ---")
    ca_clause = "3. Non-Compete. Contractor agrees not to compete with the Company anywhere in the state for 2 years."
    
    # 1. Under California jurisdiction -> MUST be critical violation
    result_ca = check_deterministic_guardrails(ca_clause, jurisdiction="California", role="Buyer")
    assert result_ca is not None, "Guardrail must trigger for CA non-compete"
    assert result_ca["risk_score"] == 10, f"Expected risk_score 10, got {result_ca['risk_score']}"
    assert "California Business and Professions Code § 16600" in result_ca["risk_analysis"][0]
    print("[PASS] California non-compete flagged with score 10 and statutory citation.")

    # 2. Under Delaware jurisdiction -> Standard non-compete rule (score 9, 6-12 month limit)
    result_de = check_deterministic_guardrails(ca_clause, jurisdiction="Delaware", role="Buyer")
    assert result_de is not None
    assert result_de["risk_score"] == 9
    assert "California" not in result_de["risk_analysis"][0]
    print("[PASS] Delaware non-compete flagged with general commercial reasonableness rule.")


def test_role_aware_liquidated_damages():
    print("--- Testing Role-Aware Liquidated Damages ---")
    ld_clause = "2. Liquidated Damages. Contractor shall pay liquidated damages of $500,000 for any project delay."

    # Seller perspective -> Defend contractor against severe delay penalties
    res_seller = check_deterministic_guardrails(ld_clause, jurisdiction="New York", role="Seller")
    assert res_seller is not None
    assert res_seller["risk_score"] == 9
    assert "Severe financial exposure for Seller/Contractor" in res_seller["risk_analysis"][0]
    print("[PASS] Seller perspective defends against uncapped delay penalties.")

    # Buyer perspective -> Evaluates adequacy / enforceable pre-estimate
    res_buyer = check_deterministic_guardrails(ld_clause, jurisdiction="New York", role="Buyer")
    assert res_buyer is not None
    assert res_buyer["risk_score"] == 7
    assert "Liquidated damages clause: Assess whether delay damages" in res_buyer["risk_analysis"][0]
    print("[PASS] Buyer perspective evaluates pre-estimate calibration.")


def test_docx_export_formatting():
    print("--- Testing Native DOCX Export Formatting ---")
    contract_mock = {
        "filename": "Master_Services_Agreement.docx",
        "role": "Seller",
        "jurisdiction": "California",
        "pii_mapping": {
            "<ORGANIZATION_1>": "Acme Corp",
            "<PERSON_1>": "Jane Doe"
        }
    }

    clauses_mock = [
        {
            "clause_index": 0,
            "original_text": "1. Parties. This Agreement is between <ORGANIZATION_1> and <PERSON_1>.",
            "risk_score": 0,
            "prosecutor_flags": [],
            "defender_redline": ""
        },
        {
            "clause_index": 1,
            "original_text": "2. Liquidated Damages. <PERSON_1> shall pay $500,000 in liquidated damages for any delay.",
            "risk_score": 9,
            "prosecutor_flags": ["Severe liquidated damages penalty"],
            "defender_redline": "2. Compensation and Penalties. Liquidated damages shall be capped at 10% of total fees paid."
        }
    ]

    buffer = build_contract_docx(contract_mock, clauses_mock)
    assert isinstance(buffer, io.BytesIO), "Must return io.BytesIO buffer"
    assert buffer.getbuffer().nbytes > 0, "Buffer must not be empty"

    # Read back generated DOCX using python-docx to verify content & formatting
    doc = docx.Document(buffer)
    full_text = "\n".join([p.text for p in doc.paragraphs])
    
    assert "CONTRACT REVIEW & REDLINE EXPORT" in full_text
    assert "Acme Corp" in full_text, "PII should be unmasked in exported docx"
    assert "Jane Doe" in full_text, "PII should be unmasked in exported docx"
    assert "Compensation and Penalties." in full_text
    assert "capped at 10%" in full_text
    
    # Check that runs have formatting applied
    has_strike = False
    has_colored = False
    for p in doc.paragraphs:
        for run in p.runs:
            if run.font.strike:
                has_strike = True
            if run.font.color and run.font.color.rgb:
                has_colored = True

    assert has_strike, "Deleted text must have strikethrough applied"
    assert has_colored, "Redline additions/deletions must have color applied"
    print("[PASS] Native DOCX generation applies strikethrough, color, and PII unmasking.")


def test_fastapi_routes():
    print("--- Testing FastAPI Route Registrations ---")
    from main import app
    route_paths = list(app.openapi()["paths"].keys())
    print(f"Registered paths: {route_paths}")

    assert "/export/docx/{contract_id}" in route_paths, "Root GET /export/docx/{contract_id} must be registered"
    assert "/api/contracts/upload" in route_paths, "Upload endpoint must be registered"
    assert "/api/contracts/analyze" in route_paths, "Analyze alias endpoint must be registered"
    assert "/api/contracts/{contract_id}/stream" in route_paths, "SSE Stream endpoint must be registered"
    assert "/api/contracts/export/docx/{contract_id}" in route_paths, "API DOCX Export endpoint must be registered"
    print("[PASS] All required REST endpoints are registered on FastAPI application.")


if __name__ == "__main__":
    test_california_non_compete_guardrail()
    test_role_aware_liquidated_damages()
    test_docx_export_formatting()
    test_fastapi_routes()
    print("\n>>> ALL TESTS PASSED SUCCESSFULLY! <<<")
